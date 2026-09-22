"""Build the published data from the student's decision tracker.

Usage: python tools/build_data.py path/to/DPI-HT-01_Master_Decision_Tracker.xlsx
"""

from __future__ import annotations

import hashlib
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT.parent / "student_package"
CODEX = PACKAGE / "04 CODEX FILES - Give These to Codex"
CASE = PACKAGE / "03 CASE FILES - Open and Investigate"
STUDENT = {"id": "pg25032", "name": "Patriks gredzens"}


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def rows(sheet, start, stop, columns):
    for row in range(start, stop + 1):
        values = [sheet.cell(row, col).value for col in range(1, columns + 1)]
        yield row, values


def main(workbook_path: Path):
    template = json.loads((CODEX / "01 GIVE TO CODEX - Answer Template.json").read_text(encoding="utf-8"))
    schema = json.loads((CODEX / "02 GIVE TO CODEX - Submission Rules.json").read_text(encoding="utf-8"))
    book = load_workbook(workbook_path, data_only=True, read_only=False)
    decisions_sheet = book["Decisions"]
    evidence_sheet = book["Evidence Index"]
    schedule_sheet = book["Supporting Schedules"]
    statements_sheet = book["Final Statements"]
    choice_sheet = book["Evidence Choices"]

    filename_by_id = {
        text(choice_sheet.cell(row, 1).value): text(choice_sheet.cell(row, 2).value)
        for row in range(2, 14)
    }
    filename_by_id[text(choice_sheet.cell(14, 1).value)] = text(choice_sheet.cell(14, 2).value)
    filename_by_number = {key[:2]: value for key, value in filename_by_id.items() if key[:2].isdigit()}
    evidence = []
    unmapped_rows = []
    for row, values in rows(evidence_sheet, 7, 125, 8):
        if not any(value not in (None, "") for value in (values[0], *values[2:])):
            continue
        evidence_id = text(values[0])
        # Excel may discard formula caches. Re-evaluate only the workbook's
        # own exact-ID lookup rather than treating a blank cache as a blank file.
        filename = text(values[1]) or filename_by_id.get(evidence_id, "")
        if evidence_id and not filename:
            unmapped_rows.append({"row": row, "evidenceId": evidence_id})
        evidence.append({
            "row": row,
            "id": evidence_id,
            "fileOrRecord": filename or None,
            "caseFileHints": [filename_by_number[code] for code in dict.fromkeys(
                re.findall(r"(?<!\d)(?:0\d|1[01])\b", evidence_id)
            ) if code in filename_by_number] if not filename else [],
            "date": text(values[2]) or None,
            "sourceType": text(values[3]) or None,
            "locationOrLink": text(values[4]) or None,
            "whatItSupports": text(values[5]) or None,
            "reliabilityOrContradictions": text(values[6]) or None,
            "decisionIds": text(values[7]) or None,
        })

    expected = {item["id"]: item for item in template["decisions"]}
    decisions = []
    for row, values in rows(decisions_sheet, 8, 107, 17):
        decision_id = text(values[0])
        if not decision_id:
            continue
        decision = {
            "id": decision_id,
            "category": text(values[1]),
            "reviewTier": text(values[2]),
            "question": text(values[3]),
            "answer": text(values[6]),
            "evidence": [text(values[13])] if text(values[13]) else [],
            "confidence": text(values[14]).lower(),
            "sourceRow": row,
        }
        if decision_id in expected and any(
            decision[key] != expected[decision_id][key]
            for key in ("category", "reviewTier", "question")
        ):
            raise ValueError(f"Template identity differs at {decision_id}; review the workbook")
        if text(values[7]):
            decision["studentReasoning"] = text(values[7])
        if text(values[4]):
            decision["aiProposal"] = text(values[4])
        if text(values[5]):
            decision["independentChallenge"] = text(values[5])
        if decision["reviewTier"] == "material_judgment":
            decision["statementEffect"] = dict(zip(
                ("profit", "cash", "assets", "liabilities", "equity"),
                (number(value) for value in values[8:13]),
            ))
        changed = text(values[15]).lower()
        if changed in ("yes", "no"):
            decision["changedFromAI"] = changed == "yes"
        decisions.append(decision)

    schedules = {}
    for row, values in rows(schedule_sheet, 8, 125, 11):
        if not text(values[0]) or not text(values[1]):
            continue
        entered = values[2:5]
        closing = sum((number(value) or 0) for value in entered[:2]) - (number(entered[2]) or 0)
        schedules.setdefault(text(values[0]), []).append({
            "row": row,
            "lineOrTransaction": text(values[1]),
            "openingBalance": number(values[2]),
            "increaseOrRecognized": number(values[3]),
            "decreaseOrSettled": number(values[4]),
            "closingBalance": closing if any(number(value) is not None for value in entered) else None,
            "profitAndLossAmount": number(values[6]),
            "cashAmount": number(values[7]),
            "decisionIds": text(values[8]) or None,
            "evidenceIds": text(values[9]) or None,
            "notesOrBasis": text(values[10]) or None,
        })

    statement_names = {
        "Profit and Loss": "profitAndLoss",
        "Cash Flow": "cashFlow",
        "Balance Sheet": "balanceSheet",
    }
    statements = {name: [] for name in statement_names.values()}
    for row, values in rows(statements_sheet, 7, 75, 6):
        group = statement_names.get(text(values[0]))
        if not group or not text(values[1]):
            continue
        statements[group].append({
            "row": row,
            "lineItem": text(values[1]),
            "amount": number(values[2]),
            "scheduleReference": text(values[3]) or None,
            "decisionIds": text(values[4]) or None,
            "evidenceIds": text(values[5]) or None,
        })

    reconciliations = []
    for row, values in rows(statements_sheet, 79, 85, 6):
        if not text(values[1]):
            continue
        reconciliations.append({
            "row": row,
            "area": text(values[0]),
            "check": text(values[1]),
            "workbookResult": text(values[2]),
            "workbookBasis": text(values[3]),
        })

    by_id = {decision["id"]: decision for decision in decisions}
    uncertainty_ids = ("D047", "D049", "D066", "D071", "D073", "D074", "D075", "D078")
    uncertainties = [{
        "decisionId": decision_id,
        "answer": by_id[decision_id]["answer"],
        "studentReasoning": by_id[decision_id].get("studentReasoning", ""),
        "confidence": by_id[decision_id]["confidence"],
        "evidence": by_id[decision_id]["evidence"],
    } for decision_id in uncertainty_ids]
    board_actions = [{
        "id": by_id[f"D{index:03d}"]["id"],
        "answer": by_id[f"D{index:03d}"]["answer"],
        "studentReasoning": by_id[f"D{index:03d}"].get("studentReasoning", ""),
        "evidence": by_id[f"D{index:03d}"]["evidence"],
        "confidence": by_id[f"D{index:03d}"]["confidence"],
    } for index in range(91, 101)]

    submission = {
        "schemaVersion": "1.0",
        "caseId": template["caseId"],
        "student": STUDENT,
        "evidence": evidence,
        "decisions": decisions,
        "schedules": schedules,
        "statements": statements,
        "reconciliations": reconciliations,
        "uncertainties": uncertainties,
        "boardRecommendation": {
            "approval": by_id["D091"]["answer"],
            "earnOut": by_id["D100"]["answer"],
            "actions": board_actions,
        },
    }

    ids = [decision["id"] for decision in decisions]
    if ids != [f"D{number:03d}" for number in range(1, 101)]:
        raise ValueError("Decisions must be D001-D100, in order, exactly once")
    if sum(decision["reviewTier"] == "material_judgment" for decision in decisions) != 25:
        raise ValueError("Expected exactly 25 material judgments")
    required = schema["required"]
    if any(key not in submission for key in required):
        raise ValueError("Submission is missing a required top-level field")
    if submission["schemaVersion"] != schema["properties"]["schemaVersion"]["const"]:
        raise ValueError("Wrong schema version")
    if submission["caseId"] != schema["properties"]["caseId"]["const"]:
        raise ValueError("Wrong case ID")
    if not all(key in submission["student"] for key in schema["properties"]["student"]["required"]):
        raise ValueError("Missing student identity")
    rules = schema["properties"]["decisions"]
    if not rules["minItems"] <= len(decisions) <= rules["maxItems"]:
        raise ValueError("Wrong decision count")
    decision_rules = rules["items"]
    for decision in decisions:
        if any(key not in decision for key in decision_rules["required"]):
            raise ValueError(f"Required decision field missing in {decision['id']}")
        if not re.fullmatch(r"D[0-9]{3}", decision["id"]):
            raise ValueError(f"Bad decision ID: {decision['id']}")
        if decision["reviewTier"] not in decision_rules["properties"]["reviewTier"]["enum"]:
            raise ValueError(f"Bad review tier: {decision['id']}")
        if decision["confidence"] not in decision_rules["properties"]["confidence"]["enum"]:
            raise ValueError(f"Bad confidence: {decision['id']}")
        if not decision["evidence"] or not all(isinstance(item, str) for item in decision["evidence"]):
            raise ValueError(f"Evidence required: {decision['id']}")
        if decision["reviewTier"] == "material_judgment":
            for key in ("aiProposal", "independentChallenge", "studentReasoning", "statementEffect", "changedFromAI"):
                if key not in decision:
                    raise ValueError(f"Material review field {key} missing: {decision['id']}")
            if len(decision["independentChallenge"]) < 20 or len(decision["studentReasoning"]) < 20:
                raise ValueError(f"Material review text too short: {decision['id']}")
            if not isinstance(decision["changedFromAI"], bool):
                raise ValueError(f"Changed-from-AI flag invalid: {decision['id']}")
            if set(decision["statementEffect"]) != {"profit", "cash", "assets", "liabilities", "equity"}:
                raise ValueError(f"Statement effect incomplete: {decision['id']}")
            if not all(value is None or isinstance(value, (int, float)) for value in decision["statementEffect"].values()):
                raise ValueError(f"Statement effect must be numeric or null: {decision['id']}")
    if not all(key in statements for key in schema["properties"]["statements"]["required"]):
        raise ValueError("One of the three statements is missing")

    amounts = {line["lineItem"]: line["amount"] for section in statements.values() for line in section}
    with (CASE / "02 Bank Export August.csv").open(encoding="utf-8-sig", newline="") as bank_file:
        bank_rows = list(csv.DictReader(bank_file))
    bank_closing_cash = int(bank_rows[-1]["Balance EUR"])
    supplier_bank_payments = sum(int(row["Debit EUR"]) for row in bank_rows if row["Reference"] in {"SUP-BOX", "SUP-GLS", "SUP-PRT", "SUP-EVT"})
    checks = [
        {"name": "Balance sheet", "difference": amounts["Total assets"] - amounts["Total liabilities and equity"]},
        {"name": "Cash roll-forward", "difference": amounts["Opening cash (1 Jan 2026)"] + amounts["Net decrease in cash"] - amounts["Closing cash (31 Aug 2026)"]},
        {"name": "Cash on balance sheet", "difference": amounts["Closing cash (31 Aug 2026)"] - amounts["Cash"]},
        {"name": "Cash versus bank export", "difference": amounts["Closing cash (31 Aug 2026)"] - bank_closing_cash},
        {"name": "Supplier cash versus bank export", "difference": -amounts["Cash paid to suppliers"] - supplier_bank_payments},
        {"name": "Equity roll-forward", "difference": amounts["Opening equity (1 Jan 2026, derived from opening balances)"] + amounts["Net profit for the period"] + amounts["Owner distributions"] - amounts["Total equity"]},
        {"name": "Revenue and receivables", "difference": 35000 + amounts["Revenue - delivered sales"] - amounts["Cash collected from customers (sales and opening receivables)"] + amounts["Bad debt write-off - customer R-17"] - amounts["Trade receivables, net of write-off"]},
        {"name": "Inventory roll-forward", "difference": 80000 + 459000 + amounts["Cost of goods sold - materials"] + amounts["Inventory write-down - damaged basement stock"] - amounts["Inventory"]},
        {"name": "PPE net", "difference": amounts["PPE, at cost"] + amounts["Accumulated depreciation"] - 191000},
        {"name": "Debt principal", "difference": 100000 + amounts["Loan advance received"] + amounts["Loan principal repaid"] - amounts["Bank loan"]},
        {"name": "Interest payable", "difference": -amounts["Interest expense"] + amounts["Cash paid for interest"] - amounts["Interest payable"]},
    ]

    diagnostics = {
        "sourceWorkbook": workbook_path.name,
        "sourceSha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "schemaValid": True,
        "decisionCount": len(decisions),
        "materialJudgmentCount": 25,
        "reconciliationChecks": [{**item, "passed": item["difference"] == 0} for item in checks],
        "reviewFlags": [
            {
                "type": "unresolved uncertainty",
                "title": "Inventory count and roll-forward differ",
                "detail": text(statements_sheet.cell(82, 4).value),
                "decisionIds": ["D075", "D091"],
            },
            {
                "type": "case evidence conflict",
                "title": "Event Things supplier balance",
                "detail": "File 06 lists EUR 59,000 unpaid for Event Things Europe and says supplier balances are independently confirmed. The bank export shows EUR 100,000 paid against the EUR 114,000 invoice, implying EUR 14,000 unpaid. The student certified EUR 14,000 in D013/D063/D084. The conflict remains visible for review; no answer has been changed. Student basis: " + by_id["D013"].get("studentReasoning", ""),
                "decisionIds": ["D013", "D063", "D084"],
            },
            {
                "type": "agent disagreement and student override",
                "title": "Approval of corrected accounts",
                "detail": by_id["D091"].get("studentReasoning", ""),
                "decisionIds": ["D091"],
            },
            {
                "type": "evidence index lookup",
                "title": "Some entered Evidence IDs have no automatic filename",
                "detail": "The workbook's Evidence Index contains labels outside its dropdown mapping. They are displayed unchanged; review these rows against the case files.",
                "rows": unmapped_rows,
            },
        ],
        "lowConfidenceDecisionIds": [decision["id"] for decision in decisions if decision["confidence"] == "low"],
        "studentOverrideIds": [decision["id"] for decision in decisions if decision.get("changedFromAI") is True],
        "agentDisagreementIds": ["D091"],
    }
    for filename, payload in (("submission.json", submission), ("diagnostics.json", diagnostics)):
        (ROOT / filename).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote 100 decisions, {len(evidence)} evidence rows, {sum(map(len, schedules.values()))} schedule rows")
    print(f"Schema valid; {sum(item['passed'] for item in diagnostics['reconciliationChecks'])}/{len(checks)} numeric checks pass")
    print(f"Evidence Index rows without mapped filenames: {len(unmapped_rows)}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]).resolve())
